"""Dump a raw PowerPoint (.pptx) into draft Markdown for render_deck.py, plus its pictures.

Every slide becomes a `## …` section with its text in reading order (top to bottom, left to right). Native
charts come out as `::: chart` blocks with their data, tables as pipe tables, pictures as image links, and
speaker notes as `::: notes`. Shapes drawn by hand (cards, bars made of rectangles) cannot be recognised: look at
the rendered slides (ppt/finalize.py --png --no-save on the raw file) and rebuild them as figures, charts or cards.

    python extract_pptx.py raw.pptx -o work/deck.md      # pictures go to work/media/
"""
import argparse, os, re, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu


def inline(paragraph):
    """Paragraph text; bold/italic are marked only where they differ within the paragraph (whole-line styling is design)."""
    runs = [r for r in paragraph.runs if r.text]
    styled = [r for r in runs if r.text.strip()]
    all_b = styled and all(r.font.bold for r in styled)
    all_i = styled and all(r.font.italic for r in styled)
    out = ''
    for r in runs:
        t = r.text
        if not t.strip(): out += t; continue
        lead, core, trail = re.match(r'^(\s*)(.*?)(\s*)$', t, re.S).groups()
        b, i = r.font.bold and not all_b, r.font.italic and not all_i
        if b and i: core = f'***{core}***'
        elif b: core = f'**{core}**'
        elif i: core = f'*{core}*'
        out += lead + core + trail
    out = re.sub(r'\*\*(\s*)\*\*', r'\1', out)      # join adjacent bold runs, keep the space between them
    return re.sub(r'[ \t]{2,}', '  ', out.replace('\x0b', ' ')).strip()


def is_bullet(paragraph):
    ppr = paragraph._p.pPr
    return ppr is not None and (ppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buChar') is not None
                                or ppr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum') is not None)


def max_size(tf):
    sizes = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else 0


def flatten(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from flatten(sh.shapes)
        else:
            yield sh


def chart_block(chart):
    plot = chart.plots[0]
    cats = [str(c) for c in plot.categories]
    series = list(plot.series)
    kind = {51: 'column', 52: 'column', 57: 'bar', 58: 'bar', 4: 'line', 65: 'line', 5: 'pie', -4120: 'pie'}.get(
        int(chart.chart_type), 'column')
    lines = [f'::: chart {{type={kind}}}', '| | ' + ' | '.join(s.name or f'Chuỗi {i + 1}' for i, s in enumerate(series)) + ' |',
             '|---|' + '---|' * len(series)]
    for i, c in enumerate(cats):
        lines.append(f'| {c} | ' + ' | '.join(fmt(s.values[i]) for s in series) + ' |')
    return '\n'.join(lines + [':::'])


def fmt(v):
    if v is None: return ''
    return (f'{v:.2f}'.rstrip('0').rstrip('.') if isinstance(v, float) else str(v))


def convert(path, media_dir, media_rel):
    prs = Presentation(path)
    sw, sh_ = prs.slide_width, prs.slide_height
    out = [f'<!-- source: {os.path.basename(path)} · {len(prs.slides)} slides · '
           f'{Emu(sw).inches:.2f} x {Emu(sh_).inches:.2f} in -->', '']
    n_img = 0
    for idx, slide in enumerate(prs.slides, 1):
        items = []
        for shp in flatten(slide.shapes):
            top, left = shp.top or 0, shp.left or 0
            if shp.has_text_frame and shp.text_frame.text.strip():
                tf = shp.text_frame
                lines = []
                for p in tf.paragraphs:
                    t = inline(p)
                    if t: lines.append(('- ' if is_bullet(p) else '') + t)
                items.append((top, left, max_size(tf), '\n'.join(lines)))
            elif getattr(shp, 'has_chart', False) and shp.has_chart:
                items.append((top, left, 0, chart_block(shp.chart)))
            elif getattr(shp, 'has_table', False) and shp.has_table:
                rows = [[inline(c.text_frame.paragraphs[0]) if len(c.text_frame.paragraphs) == 1
                         else '<br>'.join(inline(p) for p in c.text_frame.paragraphs) for c in r.cells] for r in shp.table.rows]
                n = max(len(r) for r in rows)
                tbl = ['| ' + ' | '.join(rows[0]) + ' |', '|' + '---|' * n] + ['| ' + ' | '.join(r) + ' |' for r in rows[1:]]
                items.append((top, left, 0, '\n'.join(tbl)))
            elif shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n_img += 1
                os.makedirs(media_dir, exist_ok=True)
                name = f'slide{idx:02d}_img{n_img}.{shp.image.ext}'
                open(os.path.join(media_dir, name), 'wb').write(shp.image.blob)
                full = (shp.width or 0) > 0.8 * sw and (shp.height or 0) > 0.8 * sh_
                items.append((top, left, 0, f'![]({media_rel}/{name})' + ('{.full}' if full else '')))
        # reading order: rows of roughly equal top position, then left to right
        items.sort(key=lambda t: (round(t[0] / Emu(914400 * 0.25)), t[1]))
        cands = [i for i, it in enumerate(items) if len(it[3]) >= 8 and not it[3].startswith(('![', '|', ':::'))]
        title_i = max(cands, key=lambda i: (items[i][2], -items[i][0])) if cands else None
        title = items[title_i][3].replace('\n', ' ') if title_i is not None and items[title_i][2] else f'Slide {idx}'
        out.append(f'## {title}')
        out.append(f'<!-- slide {idx} -->')
        out.append('')
        for i, it in enumerate(items):
            if i == title_i and it[2]: continue
            out += [it[3], '']
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out += ['::: notes', slide.notes_slide.notes_text_frame.text.strip(), ':::', '']
    return re.sub(r'\n{3,}', '\n\n', '\n'.join(out)).strip() + '\n', n_img


FRONT = """---
title: [Tên bài trình bày]
subtitle: [Tiêu đề phụ, hoặc xóa dòng này]
presenter: [Họ tên · đơn vị · sự kiện · ngày, hoặc xóa dòng này]
footer: [Tên ngắn ở chân slide]
cover: light
---

"""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('pptx')
    ap.add_argument('-o', '--out', help='output .md (default: next to the .pptx)')
    a = ap.parse_args()
    out = a.out or os.path.splitext(a.pptx)[0] + '.md'
    media = os.path.join(os.path.dirname(os.path.abspath(out)), 'media')
    md, n = convert(a.pptx, media, 'media')
    open(out, 'w', encoding='utf-8').write(FRONT + md)
    print(f'Wrote {out}  ({n} pictures)')


if __name__ == '__main__':
    main()
