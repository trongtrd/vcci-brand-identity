"""Render a VCCI presentation (.pptx) from structured Markdown, using brand-assets/06_Templates/PowerPoint/VCCI_Presentation.potx.

The cover and closing slide come from the front matter; every `# …` is a section divider and every `## …` a
slide. Each slide is placed on one of the template's 18 layouts; content the layouts do not cover (key-figure
cards, native charts, cards, steps, notes, tables) is drawn in the VCCI style. Markdown syntax: see
skills/vcci-brand/references/ppt-markdown.md.

    python render_deck.py deck.md -o deck_VCCI.pptx

Needs Python 3.9+, python-pptx, lxml and Pillow. Does not need PowerPoint (run finalize.py afterwards to
export a PDF and slide previews when PowerPoint is available).
"""
import argparse, copy, io, math, os, re, sys, zipfile
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'word'))
from common import find_brand_root

# ---------------------------------------------------------------- brand
BLUE, GOLD, BRONZE, BRONZE_INK = '264896', 'FFCD00', 'CC713C', '9C570C'
NAVY, MIST, INK, SLATE, LINE, IVORY, PAPER, SKY = '0B1A4F', 'E7ECF6', '161A26', '5C6478', 'E6E1D5', 'F4F0E6', 'FBF9F4', '8FA6D6'
BODY, SEMI, HEAD = 'Be Vietnam Pro', 'Be Vietnam Pro SemiBold', 'Lora Medium'
SERIES = [BLUE, SKY, BRONZE, SLATE, 'F2B18E', NAVY]   # guideline order (Bronze 20 fifth); counted from the last series; gold only for one highlighted value

LAYOUTS = {'cover': '01 Bìa sáng', 'agenda': '02 Chương trình', 'content': '03 Tiêu đề và nội dung',
           'two': '04 Hai cột', 'compare': '05 So sánh', 'dense': '06 Nội dung dày', 'chart': '07 Biểu đồ và bảng',
           'image': '08 Ảnh và chữ', 'title': '10 Chỉ có tiêu đề', 'blank': '11 Trống', 'event': '12 Bìa sự kiện',
           'section': '13 Chuyển chương', 'figures-dark': '14 Số liệu nổi bật', 'statement': '15 Tuyên bố lớn',
           'quote': '16 Trích dẫn', 'fullimage': '17 Ảnh toàn khung', 'end': '18 Kết thúc'}

X0, XW = 0.8, 11.73          # content column (inches)
Y_CONTENT, Y_BOTTOM = 2.0, 6.45
GAP = 0.25


def inch(v): return Emu(int(v * 914400))
def rgb(h): return RGBColor.from_string(h)

# ---------------------------------------------------------------- Markdown

def parse_front(text):
    meta = {}
    if text.startswith('---'):
        end = text.find('\n---', 3)
        if end != -1:
            for line in text[3:end].strip().splitlines():
                if ':' in line and not line.lstrip().startswith('#'):
                    k, v = line.split(':', 1)
                    meta[k.strip().lower()] = v.strip().strip('"').strip("'")
            text = text[end + 4:]
    return meta, text


def parse_attrs(s):
    classes, kv = set(), {}
    for m in re.finditer(r'\.([\w-]+)|([\w-]+)=("([^"]*)"|[^\s}]+)', s or ''):
        if m.group(1): classes.add(m.group(1))
        else: kv[m.group(2)] = m.group(4) if m.group(4) is not None else m.group(3)
    return classes, kv


def split_attrs(line):
    m = re.search(r'\s*\{([^{}]*)\}\s*$', line)
    if m and ('.' in m.group(1) or '=' in m.group(1)):
        return (line[:m.start()],) + parse_attrs(m.group(1))
    return line, set(), {}


def strip_comments(text):
    return re.sub(r'<!--.*?-->', '', text, flags=re.S)


def parse_slides(text):
    slides, cur = [], None
    for line in strip_comments(text).replace('\r\n', '\n').split('\n'):
        m = re.match(r'^(#{1,2})\s+(.*)$', line)
        if m and not (cur and cur.get('_fence')):
            title, cls, kv = split_attrs(m.group(2).strip())
            cur = {'kind': 'section' if m.group(1) == '#' else 'slide', 'title': title.strip(), 'cls': cls, 'kv': kv, 'lines': []}
            slides.append(cur); continue
        if cur is None: continue
        if re.match(r'^:::\s*\w', line.strip()): cur['_fence'] = True
        elif line.strip() == ':::': cur['_fence'] = False
        cur['lines'].append(line)
    for s in slides:
        s.pop('_fence', None)
        s['blocks'] = parse_blocks(s.pop('lines'))
    return slides


LIST_RE = re.compile(r'^(\s*)([-*+]|\d+[.)])\s+(.*)$')


def parse_blocks(lines):
    blocks, i = [], 0
    def special(l):
        s = l.strip()
        return not s or s.startswith(('>', '|', ':::', '![', '### ')) or LIST_RE.match(l)
    while i < len(lines):
        line, s = lines[i], lines[i].strip()
        if not s: i += 1; continue
        m = re.match(r'^:::\s*([\w-]+)\s*(\{.*\})?\s*$', s)
        if m:
            kind = m.group(1).lower(); cls, kv = parse_attrs((m.group(2) or '')[1:-1])
            body = []; i += 1
            while i < len(lines) and lines[i].strip() != ':::':
                body.append(lines[i]); i += 1
            i += 1
            blocks.append({'t': kind, 'cls': cls, 'kv': kv, 'body': body}); continue
        if s.startswith('>'):
            q = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                q.append(lines[i].strip()[1:].strip()); i += 1
            j = i
            while j < len(lines) and not lines[j].strip(): j += 1
            attrib = None
            if j < len(lines) and re.match(r'^(—|--)\s+', lines[j].strip()):
                attrib = re.sub(r'^(—|--)\s+', '', lines[j].strip()); i = j + 1
            blocks.append({'t': 'quote', 'text': ' '.join(q), 'attrib': attrib}); continue
        if s.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip().replace('\\|', '|') for c in re.split(r'(?<!\\)\|', lines[i].strip().strip('|'))]
                if not all(re.match(r'^:?-{2,}:?$', c) for c in cells if c): rows.append(cells)
                i += 1
            cls = set()
            if i < len(lines) and re.match(r'^\{[^}]*\}\s*$', lines[i].strip()):
                cls, _ = parse_attrs(lines[i].strip()[1:-1]); i += 1
            blocks.append({'t': 'table', 'rows': rows, 'cls': cls}); continue
        m = re.match(r'^!\[([^\]]*)\]\(([^)\s]+)\)(\{[^}]*\})?\s*$', s)
        if m:
            cls, kv = parse_attrs((m.group(3) or '{}')[1:-1])
            blocks.append({'t': 'image', 'caption': m.group(1), 'path': m.group(2), 'cls': cls, 'kv': kv}); i += 1; continue
        if LIST_RE.match(line):
            items = []
            while i < len(lines) and lines[i].strip() and (LIST_RE.match(lines[i]) or items):
                m = LIST_RE.match(lines[i])
                if m: items.append({'level': min(len(m.group(1).replace('\t', '    ')) // 2, 1), 'ordered': m.group(2)[0].isdigit(), 'text': m.group(3).strip()})
                else: items[-1]['text'] += ' ' + lines[i].strip()
                i += 1
            blocks.append({'t': 'list', 'items': items}); continue
        para = [s]; i += 1
        while i < len(lines) and not special(lines[i]):
            para.append(lines[i].strip()); i += 1
        blocks.append({'t': 'p', 'text': ' '.join(para)})
    return blocks


def paras_of(lines):
    """Lines inside a fenced block -> paragraphs (blank-line separated); list markers kept as bullets."""
    out, cur = [], []
    for l in lines + ['']:
        if not l.strip():
            if cur: out.append(' '.join(cur)); cur = []
        elif LIST_RE.match(l):
            if cur: out.append(' '.join(cur)); cur = []
            m = LIST_RE.match(l)
            out.append(('- ' if not m.group(2)[0].isdigit() else '1. ') + m.group(3).strip())
        else: cur.append(l.strip())
    return out


def words(text): return len(re.sub(r'[*_`]', '', text).split())

# ---------------------------------------------------------------- text helpers

INLINE = re.compile(r'(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*]))')


def add_runs(p, text, size=None, color=None, font=None, bold=None, italic=None):
    for part in INLINE.split(text):
        if not part: continue
        b, it = bold, italic
        if part.startswith('***') and part.endswith('***') and len(part) > 6: part, b, it = part[3:-3], True, True
        elif part.startswith('**') and part.endswith('**') and len(part) > 4: part, b = part[2:-2], True
        elif part.startswith('*') and part.endswith('*') and len(part) > 2: part, it = part[1:-1], True
        r = p.add_run(); r.text = part.replace('\\*', '*')
        f = r.font
        if size: f.size = Pt(size)
        if color: f.color.rgb = rgb(color)
        if font: f.name = font
        if b is not None: f.bold = b
        if it is not None: f.italic = it


def set_bullet(p, level=0, ordered=False, color=BLUE, size_pct=70):
    pPr = p._p.get_or_add_pPr()
    indent = 0.28 if level == 0 else 0.25
    pPr.set('marL', str(int((0.28 + 0.3 * level) * 914400))); pPr.set('indent', str(int(-indent * 914400)))
    for tag in ('a:buNone', 'a:buChar', 'a:buAutoNum', 'a:buClr', 'a:buSzPct', 'a:buFont'):
        for e in pPr.findall(qn(tag)): pPr.remove(e)
    clr = etree.SubElement(pPr, qn('a:buClr')); etree.SubElement(clr, qn('a:srgbClr')).set('val', color)
    etree.SubElement(pPr, qn('a:buSzPct')).set('val', str(size_pct * 1000))
    if ordered: etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')
    else: etree.SubElement(pPr, qn('a:buChar')).set('char', '■' if level == 0 else '–')


def no_bullet(p):
    pPr = p._p.get_or_add_pPr(); pPr.set('marL', '0'); pPr.set('indent', '0')
    for e in pPr.findall(qn('a:buChar')) + pPr.findall(qn('a:buAutoNum')): pPr.remove(e)
    if pPr.find(qn('a:buNone')) is None: etree.SubElement(pPr, qn('a:buNone'))


def spacing(p, after=None, before=None, line=None):
    if after is not None: p.space_after = Pt(after)
    if before is not None: p.space_before = Pt(before)
    if line is not None: p.line_spacing = line


def fill_text(tf, items, size=None, color=None, font=None, after=None, line=None, first_bold=False):
    """items: list of strings; '- x' / '1. x' / '  - x' become bullets. Clears the frame first."""
    tf.clear()
    first = True
    for raw in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        m = re.match(r'^(\s*)(- |\d+\. )(.*)$', raw)
        if m:
            add_runs(p, m.group(3), size, color, font)
            set_bullet(p, 1 if len(m.group(1)) >= 2 else 0, m.group(2).strip()[0].isdigit())
        else:
            add_runs(p, raw, size, color, font)
            no_bullet(p)
        if after is not None or line is not None: spacing(p, after=after, line=line)
    return tf


def text_box(slide, x, y, w, h, items, size=14, color=INK, font=BODY, anchor=MSO_ANCHOR.TOP, align=None, after=4, line=1.1, margin=0):
    tb = slide.shapes.add_textbox(inch(x), inch(y), inch(w), inch(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = inch(margin)
    fill_text(tf, items, size, color, font, after, line)
    if align:
        for p in tf.paragraphs: p.alignment = align
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, inch(x), inch(y), inch(w), inch(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ''
    return s


def est_lines(text, width_in, size_pt, em=0.52):
    """Rough line count; em = average glyph width in ems (Be Vietnam Pro ~0.52, Lora ~0.47)."""
    chars_per_line = max(8, width_in / (size_pt / 72 * em))
    return sum(max(1, math.ceil(len(re.sub(r'[*_]', '', t)) / chars_per_line)) for t in (text if isinstance(text, list) else [text]))


def est_height(text, width_in, size_pt, line=1.2, after_pt=4):
    items = text if isinstance(text, list) else [text]
    return est_lines(items, width_in, size_pt) * size_pt * line / 72 + len(items) * after_pt / 72

# ---------------------------------------------------------------- numbers

def num(s):
    s = s.strip().replace('−', '-').replace('−', '-')
    s = re.sub(r'[%\s]', '', s)
    if re.match(r'^[+-]?\d{1,3}(\.\d{3})+(,\d+)?$', s): s = s.replace('.', '').replace(',', '.')   # 12.500,5
    else: s = s.replace(',', '.')
    try: return float(s)
    except ValueError: return None


def decimals(cells):
    d = 0
    for c in cells:
        m = re.search(r'[.,](\d+)\s*%?$', c.strip())
        if m and not re.match(r'^\d{1,3}(\.\d{3})+$', c.strip()): d = max(d, len(m.group(1)))
    return d

# ---------------------------------------------------------------- deck builder

class Deck:
    def __init__(self, meta, base_dir, template):
        self.meta, self.base_dir, self.warnings = meta, base_dir, []
        files = zipfile.ZipFile(template)
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            for n in files.namelist():
                data = files.read(n)
                if n == '[Content_Types].xml':
                    data = data.replace(b'presentationml.template.main+xml', b'presentationml.presentation.main+xml')
                z.writestr(n, data)
        self.prs = Presentation(io.BytesIO(buf.getvalue()))
        self.layouts = {l.name: l for l in self.prs.slide_layouts}
        sl = self.prs.slides._sldIdLst          # drop the 17 sample slides
        for sid in list(sl):
            self.prs.part.drop_rel(sid.get(qn('r:id'))); sl.remove(sid)
        self.footer = meta.get('footer', meta.get('title', ''))

    # -- slide scaffolding
    def new(self, key, n):
        slide = self.prs.slides.add_slide(self.layouts[LAYOUTS[key]])
        slide._vcci_n = n
        lay = slide.slide_layout
        for sp in lay.placeholders:   # python-pptx does not copy footer / slide-number placeholders
            t = sp.placeholder_format.type
            if t is None: continue
            name = str(t)
            if 'FOOTER' in name and self.footer:
                el = copy.deepcopy(sp._element); self._set_ph_text(el, f'<a:r><a:rPr lang="vi-VN" dirty="0"/><a:t></a:t></a:r>', self.footer)
                slide.shapes._spTree.append(el)
            elif 'SLIDE_NUMBER' in name:
                el = copy.deepcopy(sp._element)
                self._set_ph_text(el, '<a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum"><a:rPr lang="vi-VN"/><a:t>‹#›</a:t></a:fld>')
                slide.shapes._spTree.append(el)
        return slide

    def _set_ph_text(self, el, inner, text=None):
        ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
        tx = el.find(qn('p:txBody'))
        for p in tx.findall(qn('a:p')): tx.remove(p)
        p = etree.fromstring(f'<a:p {ns}>{inner}</a:p>')
        if text is not None: p.find('.//' + qn('a:t')).text = text
        tx.append(p)

    def ph(self, slide, idx):
        for s in slide.placeholders:
            if s.placeholder_format.idx == idx: return s
        return None

    def drop_empty(self, slide):
        for s in list(slide.placeholders):
            if s.placeholder_format.type is not None and any(k in str(s.placeholder_format.type) for k in ('FOOTER', 'SLIDE_NUMBER')):
                continue
            if s.has_text_frame and not s.text_frame.text.strip():
                s._element.getparent().remove(s._element)

    def put(self, slide, idx, items, **kw):
        p = self.ph(slide, idx)
        if p is None: return None
        if isinstance(items, str): items = [items]
        fill_text(p.text_frame, items, **kw)
        return p

    def label(self, slide, text, dark=False):
        tb = text_box(slide, X0, 0.2, 9.5, 0.3, [text.upper()], size=11, color='C9D3EA' if dark else BRONZE_INK, font=SEMI)
        for r in tb.text_frame.paragraphs[0].runs:
            r.font._rPr.set('spc', '120')
        return tb

    def fit_title(self, slide, idx, text, base=30, width=10.0, two_line_cap=None):
        p = self.put(slide, idx, [text])
        if p is None: return
        size = base
        lines = lambda sz: est_lines(text, width, sz, em=0.47)
        if two_line_cap and lines(size) > 1:
            size = two_line_cap      # two lines must stay above the layout's gold rule
        while size > 18 and lines(size) > 2: size -= 2
        if size != base:
            for r in p.text_frame.paragraphs[0].runs: r.font.size = Pt(size)
        if lines(size) > 2: self.warn(slide, 'tiêu đề quá dài (hơn 2 dòng); nên rút gọn hoặc chuyển phần sau vào lead')

    def warn(self, slide, msg):
        self.warnings.append(f'Slide {slide._vcci_n}: {msg}')

    def notes(self, slide, blocks):
        for b in blocks:
            if b['t'] == 'notes':
                ns = slide.notes_slide
                if ns.notes_text_frame is None:   # the template's notes master has no body placeholder
                    sp = parse_xml(
                        '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:nvSpPr><p:cNvPr id="3" name="Notes"/>'
                        '<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="body" idx="1"/></p:nvPr></p:nvSpPr>'
                        '<p:spPr><a:xfrm><a:off x="685800" y="4400550"/><a:ext cx="5486400" cy="3600450"/></a:xfrm></p:spPr>'
                        '<p:txBody><a:bodyPr/><a:lstStyle/><a:p/></p:txBody></p:sp>')
                    ns.shapes._spTree.append(sp)
                ns.notes_text_frame.text = '\n'.join(l.rstrip() for l in b['body']).strip()

    def picture_fill(self, slide, path, x, y, w, h):
        """Picture cropped to fill the box (for photo frames)."""
        from PIL import Image
        full = self.path(path)
        if not full: return None
        with Image.open(full) as im: iw, ih = im.size
        pic = slide.shapes.add_picture(full, inch(x), inch(y), inch(w), inch(h))
        r_img, r_box = iw / ih, w / h
        if r_img > r_box:
            c = (1 - r_box / r_img) / 2; pic.crop_left = pic.crop_right = c
        else:
            c = (1 - r_img / r_box) / 2; pic.crop_top = pic.crop_bottom = c
        return pic

    def picture_fit(self, slide, path, x, y, w, h):
        from PIL import Image
        full = self.path(path)
        if not full: return None
        with Image.open(full) as im: iw, ih = im.size
        s = min(w / iw, h / ih); pw, ph = iw * s, ih * s
        return slide.shapes.add_picture(full, inch(x + (w - pw) / 2), inch(y + (h - ph) / 2), inch(pw), inch(ph))

    def path(self, p):
        full = p if os.path.isabs(p) else os.path.join(self.base_dir, p)
        if not os.path.exists(full):
            self.warnings.append(f'Image not found: {p}'); return None
        return full

    def replace_ph_with_picture(self, slide, idx, path, mode='fill'):
        ph = self.ph(slide, idx)
        if ph is None or not path: return
        x, y, w, h = (Emu(v).inches for v in (ph.left, ph.top, ph.width, ph.height))
        ph._element.getparent().remove(ph._element)
        (self.picture_fill if mode == 'fill' else self.picture_fit)(slide, path, x, y, w, h)

    def fade(self, slide):
        el = slide._element
        tr = etree.SubElement(el, qn('p:transition')); tr.set('spd', 'med'); etree.SubElement(tr, qn('p:fade'))
        # p:transition must precede p:timing / p:extLst
        for tag in ('p:timing', 'p:extLst'):
            for e in el.findall(qn(tag)): el.remove(e); el.append(e)

    # -- fixed slides
    def cover(self, n):
        key = 'event' if self.meta.get('cover', 'light').lower() in ('event', 'dark', 'blue') else 'cover'
        s = self.new(key, n)
        self.fit_title(s, 102, self.meta.get('title', ''), base=34 if key == 'cover' else 36, width=6.9 if key == 'cover' else 7.6)
        self.put(s, 103, [self.meta['subtitle']]) if self.meta.get('subtitle') else None
        self.put(s, 104, [self.meta['presenter']]) if self.meta.get('presenter') else None
        self.drop_empty(s)

    def end(self, n):
        s = self.new('end', n)
        self.put(s, 100, [self.meta.get('end_title', 'Cảm ơn')])
        self.put(s, 101, [self.meta.get('end_subtitle', 'Thank you')])
        contact = [l.strip() for l in self.meta.get('contact', '').split('|') if l.strip()]
        self.put(s, 103, ['**Liên đoàn Thương mại và Công nghiệp Việt Nam**', 'Số 9 Đào Duy Anh, Đống Đa, Hà Nội  ·  vcci.com.vn'] + contact)
        self.drop_empty(s)

    # -- slides
    def section(self, sd, n, number):
        s = self.new('section', n)
        self.put(s, 100, [sd['kv'].get('number', f'{number:02d}')])
        self.fit_title(s, 101, sd['title'], base=40, width=7.6)
        sub = next((b['text'] for b in sd['blocks'] if b['t'] == 'p'), None)
        if sub: self.put(s, 102, [sub])
        icon = sd['kv'].get('icon')
        if icon and icon.lower() != 'none': self.replace_ph_with_picture(s, 103, hero_icon(icon), 'fit')
        self.drop_empty(s); self.fade(s); self.notes(s, sd['blocks'])

    def choose(self, sd):
        if sd['kv'].get('layout'): return sd['kv']['layout']
        t = [b['t'] for b in sd['blocks'] if b['t'] not in ('notes', 'source', 'lead')]
        if 'statement' in sd['cls']: return 'statement'
        if t == ['quote']: return 'quote'
        if 'agenda' in t: return 'agenda'
        if t.count('col') == 2 and set(t) <= {'col'}: return 'two'
        if set(t) <= {'left', 'right'} and t: return 'compare'
        imgs = [b for b in sd['blocks'] if b['t'] == 'image']
        if imgs and 'full' in imgs[0]['cls'] and set(t) <= {'image'}: return 'fullimage'
        if imgs and set(t) <= {'image', 'p', 'list'} and 'photo' in imgs[0]['cls']: return 'image'
        if 'figures' in t and 'dark' in sd['cls'] and set(t) <= {'figures'}: return 'figures-dark'
        if 'chart' in t and t.count('chart') == 1 and set(t) <= {'chart', 'p', 'list', 'figures'} and 'wide' not in sd['cls']: return 'chart'
        if set(t) <= {'p', 'list'} and t:
            w = sum(words(b.get('text', '')) + sum(words(i['text']) for i in b.get('items', [])) for b in sd['blocks'])
            return 'dense' if (w > 90 or 'dense' in sd['cls']) else 'content'
        return 'title'

    def text_items(self, blocks):
        out = []
        for b in blocks:
            if b['t'] == 'p': out.append(b['text'])
            elif b['t'] == 'list':
                out += [('  ' if it['level'] else '') + ('1. ' if it['ordered'] else '- ') + it['text'] for it in b['items']]
        return out

    def slide(self, sd, n):
        key = self.choose(sd)
        if key not in LAYOUTS: self.warnings.append(f'Slide {n}: unknown layout "{key}"'); key = 'title'
        s = self.new(key, n)
        blocks = sd['blocks']
        lead = next((b for b in blocks if b['t'] == 'lead'), None)
        source = next((b for b in blocks if b['t'] == 'source'), None)
        dark = key in ('figures-dark', 'statement', 'quote', 'fullimage')
        if sd['kv'].get('label') and key not in ('statement', 'quote', 'fullimage'): self.label(s, sd['kv']['label'], dark)

        if key == 'statement':
            self.fit_title(s, 101, sd['title'], base=40, width=10.4)
            rest = self.text_items(blocks)
            if rest: self.put(s, 102, rest)
        elif key == 'quote':
            q = next(b for b in blocks if b['t'] == 'quote')
            self.put(s, 101, [f'“{q["text"].strip(chr(34) + "“”")}”'])
            if q['attrib']: self.put(s, 102, [q['attrib']])
        elif key == 'figures-dark':
            self.fit_title(s, 100, sd['title'], base=30, width=11.7)
            items = self.figure_items(next(b for b in blocks if b['t'] == 'figures'))[:3]
            for k, (v, l, note) in enumerate(items):
                self.put(s, 102 + 3 * k, [v]); self.put(s, 103 + 3 * k, [l] + ([note] if note else []))
            first = self.ph(s, 102)
            if first is not None:
                for r in first.text_frame.paragraphs[0].runs: r.font.color.rgb = rgb(GOLD)
        else:
            self.fit_title(s, 103 if key == 'image' else 101, sd['title'], base=30, width=5.6 if key == 'image' else 10.0,
                           two_line_cap=None if key == 'image' else 24)
            top = Y_CONTENT
            if lead:
                lt = ' '.join(paras_of(lead['body']))
                h = est_height(lt, XW, 15, 1.15) + 0.05
                text_box(s, X0, 1.72, XW, h, [lt], size=15, color=SLATE, after=0)
                top = max(Y_CONTENT, 1.72 + h + 0.2)
            bottom = Y_BOTTOM
            if source and key != 'chart':
                text_box(s, X0, 6.45, XW, 0.3, [' '.join(paras_of(source['body']))], size=10, color=SLATE, anchor=MSO_ANCHOR.BOTTOM, after=0)
                bottom = 6.3
            self.body(s, key, sd, blocks, top, bottom, source)
        self.drop_empty(s); self.notes(s, blocks)

    def shift_ph(self, ph, top):
        """Move a placeholder's top edge down to `top`, keeping its other edges (inherited from the layout)."""
        if ph is None: return
        x, y, w, h = ph.left, ph.top, ph.width, ph.height     # inherited values, read before writing any
        if top > Emu(y).inches:
            ph.left, ph.width = x, w
            ph.top, ph.height = inch(top), Emu(int(h - (inch(top) - y)))

    def body(self, s, key, sd, blocks, top, bottom, source):
        content = [b for b in blocks if b['t'] not in ('notes', 'source', 'lead')]
        if key in ('content', 'dense'):
            items = self.text_items(content)
            if key == 'content':
                p = self.put(s, 103, items); self.shift_ph(p, top)
                if est_height(items, XW, 18, 1.1, 8) > bottom - top: self.warn(s, 'nhiều chữ so với cỡ 18 pt; tách slide hoặc dùng {layout=dense}')
            else:
                half = self.split_half(items)
                for idx, part in ((103, half[0]), (104, half[1])):
                    p = self.put(s, idx, part, size=14); self.shift_ph(p, top)
        elif key == 'two':
            cols = [b for b in content if b['t'] == 'col']
            for idx, c in zip((103, 104), cols):
                p = self.put(s, idx, paras_of(c['body'])); self.shift_ph(p, top)
        elif key == 'compare':
            for (hidx, bidx), side in zip(((103, 105), (106, 108)), ('left', 'right')):
                b = next((x for x in content if x['t'] == side), None)
                if b is None: continue
                self.put(s, hidx, [b['kv'].get('title', '')]); self.put(s, bidx, paras_of(b['body']))
        elif key == 'agenda':
            b = next(x for x in content if x['t'] == 'agenda')
            p = self.put(s, 103, paras_of(b['body'])); self.shift_ph(p, top)
        elif key == 'chart':
            ch = next(b for b in content if b['t'] == 'chart')
            self.chart(s, ch, X0, top, 8.03, (6.15 if not source else 6.1) - top)
            msg = self.text_items([b for b in content if b['t'] in ('p', 'list')])
            for f in [b for b in content if b['t'] == 'figures']:
                for v, l, note in self.figure_items(f):
                    msg.append(f'**{v}** {l}' + (f' ({note})' if note else ''))
            p = self.put(s, 103, msg, size=15 if sum(map(words, msg)) > 30 else None); self.shift_ph(p, top)
            if source: self.put(s, 104, [' '.join(paras_of(source['body']))])
            if est_height(msg, 3.3, 15, 1.1, 8) > 6.1 - top: self.warn(s, 'phần chữ bên phải biểu đồ quá dài')
        elif key == 'image':
            img = next(b for b in content if b['t'] == 'image')
            self.replace_ph_with_picture(s, 101, img['path'], 'fill')
            self.put(s, 105, self.text_items(content))
        elif key == 'fullimage':
            img = next(b for b in content if b['t'] == 'image')
            self.replace_ph_with_picture(s, 101, img['path'], 'fit' if 'fit' in img['cls'] else 'fill')
            self.put(s, 103, [img['caption'] or sd['title']])
        else:
            self.stack(s, content, top, bottom)

    def split_half(self, items):
        total = sum(len(i) for i in items); acc = 0
        for k, it in enumerate(items):
            acc += len(it)
            if acc >= total / 2: return items[:k + 1], items[k + 1:]
        return items, []

    # -- free composition on "10 Chỉ có tiêu đề"
    def stack(self, s, content, top, bottom):
        fixed, flex = [], []
        for b in content:
            h = self.fixed_height(b)
            (flex if h is None else fixed).append(b)
            b['_h'] = h
        need = sum(b['_h'] for b in fixed) + GAP * (len(content) - 1)
        room = bottom - top - need
        min_flex = sum(self.min_height(b) for b in flex)
        if room < min_flex - 0.05:
            self.warn(s, f'nội dung tràn khoảng {min_flex - room:.1f} in; tách slide hoặc bớt thành phần')
        share = max(room, min_flex) / len(flex) if flex else 0
        for b in flex: b['_h'] = max(self.min_height(b), share)
        y = top
        for b in content:
            self.draw(s, b, X0, y, XW, b['_h'])
            y += b['_h'] + GAP

    def min_height(self, b):
        return {'chart': 2.2, 'table': 0.4 * (len(b.get('rows', [])) + 0.5), 'image': 1.5}.get(b['t'], self.cards_height(b) if b['t'] == 'cards' else 1.0)

    def fixed_height(self, b):
        t = b['t']
        if t == 'p': return est_height(b['text'], XW, 16, 1.15, 6)
        if t == 'list': return est_height([i['text'] for i in b['items']], XW - 0.3, 16, 1.15, 6)
        if t == 'figures':
            items = self.figure_items(b); w = (XW - GAP * (len(items) - 1)) / max(1, len(items))
            return 0.25 + 0.75 + max(est_height(l, w - 0.4, 13, 1.2, 4) + (0.3 if n else 0) for _, l, n in items) + 0.3
        if t == 'note': return est_height(paras_of(b['body']), XW - 0.5, 14, 1.15, 4) + 0.3
        if t == 'steps':
            items = self.step_items(b); w = (XW - 0.45 * (len(items) - 1)) / max(1, len(items))
            return (0.35 if b['kv'].get('title') else 0) + 0.3 + max(est_height(tt, w - 0.3, 14, 1.1, 2) + (est_height(d, w - 0.3, 12, 1.1, 0) if d else 0) for tt, d in items) + 0.15
        if t == 'quote': return est_height(b['text'], XW, 20, 1.2, 0) + (0.4 if b['attrib'] else 0)
        if t == 'cards' and 'fill' not in b['cls']: return self.cards_height(b)
        if t in ('chart', 'table', 'image', 'cards'): return None
        return 0.5

    def draw(self, s, b, x, y, w, h):
        t = b['t']
        if t in ('p', 'list'):
            text_box(s, x, y, w, h, self.text_items([b]), size=16, after=6)
        elif t == 'figures': self.figures(s, b, x, y, w, h)
        elif t == 'note': self.note(s, b, x, y, w, h)
        elif t == 'steps': self.steps(s, b, x, y, w, h)
        elif t == 'cards': self.cards(s, b, x, y, w, h)
        elif t == 'chart': self.chart(s, b, x, y, w, h)
        elif t == 'table': self.table(s, b, x, y, w, h)
        elif t == 'image': self.picture_fit(s, b['path'], x, y, w, h)
        elif t == 'quote':
            text_box(s, x, y, w, h, [f'“{b["text"]}”'] + ([f'— {b["attrib"]}'] if b['attrib'] else []), size=20, color=BLUE, font='Lora', after=6)
        else:
            self.warnings.append(f'Unknown block ::: {t}')

    # -- components
    def figure_items(self, b):
        out = []
        for l in b['body']:
            if not l.strip(): continue
            parts = [p.strip() for p in l.split('|')] + ['', '']
            out.append((parts[0], parts[1], parts[2]))
        return out

    def figures(self, s, b, x, y, w, h):
        items = self.figure_items(b)[:4]
        n = len(items); cw = (w - GAP * (n - 1)) / n
        for k, (v, l, note) in enumerate(items):
            cx = x + k * (cw + GAP)
            rect(s, cx, y, cw, h, MIST)
            rect(s, cx, y, 0.06, h, BLUE)
            text_box(s, cx + 0.25, y + 0.2, cw - 0.4, 0.75, [v], size=36, color=BLUE, font=HEAD, after=0, line=1.0)
            text_box(s, cx + 0.25, y + 0.98, cw - 0.4, h - 1.0, [l] + ([note] if note else []), size=13, after=4)
            if note:   # comparison line in slate
                tb = s.shapes[-1].text_frame.paragraphs[-1]
                for r in tb.runs: r.font.color.rgb = rgb(SLATE); r.font.size = Pt(11)

    def note(self, s, b, x, y, w, h):
        rect(s, x, y, w, h, MIST); rect(s, x, y, 0.06, h, BLUE)
        text_box(s, x + 0.3, y + 0.15, w - 0.5, h - 0.3, paras_of(b['body']), size=14, after=4, anchor=MSO_ANCHOR.MIDDLE)

    def step_items(self, b):
        return [((l.split('|', 1) + [''])[0].strip(), (l.split('|', 1) + [''])[1].strip()) for l in b['body'] if l.strip()]

    def steps(self, s, b, x, y, w, h):
        items = self.step_items(b); n = len(items); arrow = 0.45
        if b['kv'].get('title'):
            text_box(s, x, y, w, 0.3, [b['kv']['title']], size=13, color=SLATE, font=SEMI, after=0)
            y += 0.35; h -= 0.35
        cw = (w - arrow * (n - 1)) / n
        for k, (title, desc) in enumerate(items):
            cx = x + k * (cw + arrow)
            rect(s, cx, y, cw, h, MIST)
            text_box(s, cx + 0.15, y + 0.15, cw - 0.3, h - 0.3, [f'**{title}**'] + ([desc] if desc else []), size=14, color=INK, after=2)
            tf = s.shapes[-1].text_frame
            for r in tf.paragraphs[0].runs: r.font.color.rgb = rgb(BLUE)
            for p in tf.paragraphs[1:]:
                for r in p.runs: r.font.size = Pt(12)
            if k < n - 1:
                a = s.shapes.add_shape(MSO_SHAPE.CHEVRON, inch(cx + cw + 0.12), inch(y + h / 2 - 0.16), inch(0.22), inch(0.32))
                a.fill.solid(); a.fill.fore_color.rgb = rgb(BRONZE); a.line.fill.background(); a.shadow.inherit = False

    def card_items(self, b):
        cards, cur = [], None
        for l in b['body']:
            m = re.match(r'^###\s+(.*)$', l.strip())
            if m: cur = {'title': m.group(1), 'lines': []}; cards.append(cur)
            elif cur is not None: cur['lines'].append(l)
        return [(c['title'], paras_of(c['lines'])) for c in cards]

    def cards_height(self, b):
        items = self.card_items(b); n = max(1, len(items)); cw = (XW - GAP * (n - 1)) / n
        return max(0.3 + est_height(t, cw - 0.4, 17, 1.1, 0) + 0.15 + est_height(body, cw - 0.6, 14, 1.15, 4) + 0.2 for t, body in items) if items else 1.0

    def cards(self, s, b, x, y, w, h):
        items = self.card_items(b); n = len(items); cw = (w - GAP * (n - 1)) / n
        for k, (title, body) in enumerate(items):
            cx = x + k * (cw + GAP)
            rect(s, cx, y, cw, h, MIST); rect(s, cx, y, cw, 0.06, BLUE)
            th = est_height(title, cw - 0.4, 17, 1.1, 0)
            text_box(s, cx + 0.22, y + 0.25, cw - 0.4, th + 0.05, [title], size=17, color=BLUE, font=HEAD, after=0, line=1.05)
            text_box(s, cx + 0.22, y + 0.3 + th + 0.12, cw - 0.4, h - th - 0.6, body, size=14, after=4)

    def table(self, s, b, x, y, w, h):
        rows = [r for r in b['rows'] if r]
        ncol = max(len(r) for r in rows); rows = [r + [''] * (ncol - len(r)) for r in rows]
        size = 14 if len(rows) <= 6 else 12
        row_h = max(0.36, min(0.6, h / len(rows)))
        lens = [max(4, min(45, max(len(re.sub(r'[*_]', '', r[c])) for r in rows))) for c in range(ncol)]
        widths = [w * l / sum(lens) for l in lens]
        numeric = [all(num(r[c]) is not None or not r[c] for r in rows[1:]) and any(r[c] for r in rows[1:]) for c in range(ncol)]
        gf = s.shapes.add_table(len(rows), ncol, inch(x), inch(y), inch(w), inch(row_h * len(rows)))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for k in list(tblPr): tblPr.remove(k)
        for a in list(tblPr.attrib): del tblPr.attrib[a]
        for c, cw in enumerate(widths): tbl.columns[c].width = inch(cw)
        total = 'total' in b['cls']
        for ri, r in enumerate(rows):
            tbl.rows[ri].height = inch(row_h)
            for c, txt in enumerate(r):
                cell = tbl.cell(ri, c)
                cell.margin_left = cell.margin_right = inch(0.12); cell.margin_top = cell.margin_bottom = inch(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                head = ri == 0; tot = total and ri == len(rows) - 1
                fill_text(cell.text_frame, [x.strip() for x in re.split(r'<br\s*/?>', txt)] or [''], size=size,
                          color='FFFFFF' if head else INK, font=SEMI if (head or tot) else BODY, after=0)
                for p in cell.text_frame.paragraphs:
                    p.alignment = PP_ALIGN.RIGHT if numeric[c] else PP_ALIGN.LEFT
                tcPr = cell._tc.get_or_add_tcPr()
                fill = BLUE if head else (IVORY if tot else None)
                border = (f'<a:lnL w="0"><a:noFill/></a:lnL><a:lnR w="0"><a:noFill/></a:lnR><a:lnT w="0"><a:noFill/></a:lnT>'
                          f'<a:lnB w="{19050 if (head or tot) else 9525}"><a:solidFill><a:srgbClr val="{BLUE if tot else LINE}"/></a:solidFill></a:lnB>')
                bg = f'<a:solidFill><a:srgbClr val="{fill}"/></a:solidFill>' if fill else '<a:noFill/>'
                ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
                for e in etree.fromstring(f'<x {ns}>{border}{bg}</x>'): tcPr.append(e)
        if row_h * len(rows) > h + 0.05: self.warn(s, 'bảng cao hơn chỗ trống; tách bảng hoặc bớt dòng')

    def chart(self, s, b, x, y, w, h):
        rows = [[c.strip() for c in re.split(r'(?<!\\)\|', l.strip().strip('|'))] for l in b['body'] if l.strip().startswith('|')]
        rows = [r for r in rows if not all(re.match(r'^:?-{2,}:?$', c) for c in r if c)]
        if len(rows) < 2: self.warnings.append('Chart without data'); return
        head, data = rows[0], rows[1:]
        kind = b['kv'].get('type', 'column')
        cats = [r[0] for r in data]
        cd = CategoryChartData(); cd.categories = cats
        cells = [r[k] for r in data for k in range(1, len(head))]
        dec = int(b['kv']['decimals']) if 'decimals' in b['kv'] else decimals(cells)
        pct = any('%' in c for c in cells) or b['kv'].get('unit') == '%'
        fmt = ('0' + ('.' + '0' * dec if dec else '')) + ('"%"' if pct else '')
        names = head[1:] if len(head) > 1 else ['']
        for k, name in enumerate(names):
            cd.add_series(name or ' ', [num(r[k + 1]) if k + 1 < len(r) else None for r in data], number_format=fmt)
        ct = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'barh': XL_CHART_TYPE.BAR_CLUSTERED,
              'line': XL_CHART_TYPE.LINE_MARKERS, 'stacked': XL_CHART_TYPE.COLUMN_STACKED, 'pie': XL_CHART_TYPE.DOUGHNUT}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        gf = s.shapes.add_chart(ct, inch(x), inch(y), inch(w), inch(h), cd)
        ch = gf.chart
        ch.font.name = BODY; ch.font.size = Pt(12); ch.font.color.rgb = rgb(INK)
        multi = len(names) > 1 and any(names)
        ch.has_legend = multi or kind == 'pie'
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM if kind != 'pie' else XL_LEGEND_POSITION.RIGHT
            ch.legend.include_in_layout = False; ch.legend.font.size = Pt(12)
        if b['kv'].get('title'):
            ch.has_title = True; ch.chart_title.text_frame.text = b['kv']['title']
            for r in ch.chart_title.text_frame.paragraphs[0].runs: r.font.size = Pt(13); r.font.name = SEMI; r.font.color.rgb = rgb(INK)
        else: ch.has_title = False
        plot = ch.plots[0]
        highlight = [h_.strip() for h_ in b['kv'].get('highlight', '').split(',') if h_.strip()]
        signed = 'signed' in b['cls']
        if kind != 'pie':
            plot.gap_width = 55
            if hasattr(plot, 'overlap') and multi: plot.overlap = -8
            va, ca = ch.value_axis, ch.category_axis
            va.has_major_gridlines = False; va.visible = False
            ca.format.line.color.rgb = rgb(LINE); ca.has_major_gridlines = False
            ca.tick_labels.font.size = Pt(12); ca.tick_labels.font.color.rgb = rgb(INK)
            if kind in ('bar', 'barh'): ca.reverse_order = True
            vals = [v for r in data for v in [num(c) for c in r[1:]] if v is not None]
            if vals and min(vals) >= 0 and 'zero' not in b['cls']: va.minimum_scale = 0
            plot.has_data_labels = True
            dl = plot.data_labels; dl.font.size = Pt(12); dl.font.color.rgb = rgb(INK); dl.number_format = fmt; dl.number_format_is_linked = False
            if kind not in ('line', 'stacked'): dl.position = XL_LABEL_POSITION.OUTSIDE_END
        else:
            plot.has_data_labels = True; plot.data_labels.number_format = fmt; plot.data_labels.number_format_is_linked = False
        for k, ser in enumerate(plot.series):
            # the last series (usually the latest year) is the darkest; earlier series step back
            color = SERIES[(len(names) - 1 - k) % len(SERIES)] if multi else BLUE
            if kind == 'line':
                ser.format.line.color.rgb = rgb(color); ser.format.line.width = Pt(2.5); ser.smooth = False
            elif kind != 'pie':
                ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(color)
                ser.invert_if_negative = False
            for i, c in enumerate(cats):
                v = num(data[i][k + 1]) if k + 1 < len(data[i]) else None
                pt_color = None
                if kind == 'pie': pt_color = SERIES[i % len(SERIES)]
                elif signed and v is not None and v < 0: pt_color = BRONZE
                if c in highlight and (not multi or k == len(names) - 1): pt_color = GOLD
                if pt_color:
                    pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = rgb(pt_color)
                    _no_invert(ser, i)
                raw = data[i][k + 1].strip() if k + 1 < len(data[i]) else ''
                if raw and kind != 'pie' and 'auto-labels' not in b['cls']:
                    # label shows the number exactly as written in the source (decimal comma, sign, %)
                    if pct and not raw.endswith('%'): raw += '%'
                    tf = ser.points[i].data_label.text_frame
                    tf.text = raw
                    for r in tf.paragraphs[0].runs: r.font.size = Pt(12); r.font.color.rgb = rgb(INK); r.font.name = BODY
                    if kind not in ('line', 'stacked'): ser.points[i].data_label.position = XL_LABEL_POSITION.OUTSIDE_END
        if signed and kind in ('bar', 'barh', 'column'):
            # category labels at the low end so they do not collide with negative bars
            ca = ch.category_axis._element
            for e in ca.findall(qn('c:tickLblPos')): e.set('val', 'low')
        if sum(1 for c in cats for k in range(len(names)) if c in highlight) > 1:
            self.warnings.append('Chart: more than one gold highlight; keep one gold accent per slide')
        return gf


def _no_invert(ser, i):
    """Data points need their own invertIfNegative=0, or PowerPoint draws negative bars hollow."""
    for dpt in ser._element.findall(qn('c:dPt')):
        if dpt.find(qn('c:idx')).get('val') == str(i) and dpt.find(qn('c:invertIfNegative')) is None:
            e = etree.Element(qn('c:invertIfNegative')); e.set('val', '0')
            dpt.find(qn('c:idx')).addnext(e)


def hero_icon(spec):
    folder = os.path.join(find_brand_root(), 'brand-assets', '05_Icons', 'Hero', 'PNG', 'Gold-on-dark')
    key = spec.strip().lower()
    for f in sorted(os.listdir(folder)):
        m = re.match(r'VCCI_Hero_(\d\d)_(.+)_512px\.png$', f)
        if m and key in (m.group(1), m.group(1).lstrip('0'), m.group(2)):
            return os.path.join(folder, f)
    raise SystemExit(f'Unknown icon "{spec}". Use a number 1-24 or a slug from brand-assets/05_Icons/Hero.')


def render(md_path, out_path, template=None):
    text = open(md_path, encoding='utf-8-sig').read()
    meta, body = parse_front(text)
    slides = parse_slides(body)
    template = template or os.path.join(find_brand_root(), 'brand-assets', '06_Templates', 'PowerPoint', 'VCCI_Presentation.potx')
    d = Deck(meta, os.path.dirname(os.path.abspath(md_path)), template)
    n = 1
    if meta.get('cover', 'light').lower() != 'none': d.cover(n); n += 1
    sec = 0
    for sd in slides:
        if sd['kind'] == 'section': sec += 1; d.section(sd, n, sec)
        else: d.slide(sd, n)
        n += 1
    if meta.get('end', 'true').lower() not in ('false', 'no', 'none'): d.end(n); n += 1
    d.prs.save(out_path)
    return d.warnings, n - 1


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('markdown')
    ap.add_argument('-o', '--out')
    ap.add_argument('--template')
    a = ap.parse_args()
    out = a.out or os.path.splitext(a.markdown)[0] + '_VCCI.pptx'
    warnings, n = render(a.markdown, out, a.template)
    for w in warnings: print('WARNING:', w)
    print(f'Wrote {out}  ({n} slides)')


if __name__ == '__main__':
    main()
